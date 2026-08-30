from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from scripts.util.write_ensemble_optimizer_report import (
    APPROACHES,
    LEVELS,
    LEVEL_METRICS,
    OVERALL_METRICS,
    write_report,
)


def _metrics_frame(include_folds: bool) -> pd.DataFrame:
    metric_names = OVERALL_METRICS.copy()
    for level in LEVELS:
        metric_names.extend([f"{level}_Support"] + [f"{level}_{metric}" for metric in LEVEL_METRICS])
    rows = []
    folds = ["1", "2", "Average"] if include_folds else [None]
    for fold in folds:
        for approach_index, approach in enumerate(APPROACHES):
            for metric_index, metric in enumerate(metric_names):
                row = {"Approach": approach, "Metric": metric, "Value": float(approach_index + metric_index + 1)}
                if include_folds:
                    row["Fold"] = fold
                rows.append(row)
    return pd.DataFrame(rows)


def test_write_report_uses_julia_metric_tables(tmp_path: Path) -> None:
    _metrics_frame(include_folds=False).to_csv(tmp_path / "metrics_summary.csv", index=False)
    _metrics_frame(include_folds=True).to_csv(tmp_path / "fold_metrics_summary.csv", index=False)
    pd.DataFrame(
        {
            "Fold": [1, 2],
            "Grouping": ["All", "All"],
            "Model_Slug": ["model-a", "model-a"],
            "Weight": [0.4, 0.6],
        }
    ).to_csv(tmp_path / "fold_weights.csv", index=False)

    workbook_path = write_report(tmp_path, remove_source_csvs=True)

    workbook = load_workbook(workbook_path, read_only=True)
    assert workbook.sheetnames == [
        "Overall Metrics",
        "Per-Level Metrics",
        "Overall Metrics by Fold",
        "Per-Level Metrics by Fold",
        "Fold Weight Ranges",
    ]
    assert [row[0] for row in workbook["Overall Metrics by Fold"].iter_rows(min_row=2, values_only=True)] == ["1"] * 7 + ["2"] * 7 + ["Average"] * 7
    assert not (tmp_path / "metrics_summary.csv").exists()
    assert not (tmp_path / "fold_metrics_summary.csv").exists()
    assert not (tmp_path / "fold_weights.csv").exists()