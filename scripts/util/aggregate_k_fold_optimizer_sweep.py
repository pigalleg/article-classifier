"""Build a cross-k comparison workbook from completed optimizer reports."""

from __future__ import annotations

import argparse
import tomllib
from copy import copy
from pathlib import Path

import pandas as pd


REQUIRED_SHEETS = ["Overall Metrics", "Per-Level Metrics", "Fold Weight Ranges"]


def completed_child_directories(sweep_dir: Path) -> list[tuple[int, Path]]:
    child_directories: list[tuple[int, Path]] = []
    for child_dir in sweep_dir.iterdir():
        if not child_dir.is_dir() or not child_dir.name.startswith("k"):
            continue
        try:
            fold_count = int(child_dir.name[1:])
        except ValueError:
            continue
        required_paths = [
            child_dir / "optimization_report.xlsx",
            child_dir / "run_metadata.toml",
            child_dir / "final_weights.csv",
        ]
        missing_paths = [path.name for path in required_paths if not path.is_file()]
        if missing_paths:
            raise FileNotFoundError(
                f"Incomplete child run {child_dir}: missing {', '.join(missing_paths)}."
            )
        child_directories.append((fold_count, child_dir))
    if not child_directories:
        raise FileNotFoundError(f"No completed k-fold child runs found in {sweep_dir}.")
    return sorted(child_directories)


def _read_sheet(report_path: Path, sheet_name: str, fold_count: int) -> pd.DataFrame:
    try:
        frame = pd.read_excel(report_path, sheet_name=sheet_name)
    except ValueError as error:
        raise ValueError(f"{report_path} is missing required sheet {sheet_name!r}.") from error
    frame.insert(0, "Fold_Count", fold_count)
    return frame


def _read_metadata(metadata_path: Path, fold_count: int) -> pd.DataFrame:
    with metadata_path.open("rb") as metadata_file:
        metadata = tomllib.load(metadata_file)
    metadata["Fold_Count"] = fold_count
    return pd.DataFrame([metadata])


def _format_workbook(workbook_path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            font = copy(cell.font)
            font.bold = True
            cell.font = font
        for column_cells in worksheet.columns:
            column_letter = column_cells[0].column_letter
            width = max(len(str(cell.value or "")) for cell in column_cells) + 2
            worksheet.column_dimensions[column_letter].width = min(width, 38)
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
    workbook.save(workbook_path)


def write_comparison_workbook(sweep_dir: Path) -> Path:
    overall_metrics: list[pd.DataFrame] = []
    level_metrics: list[pd.DataFrame] = []
    weight_stability: list[pd.DataFrame] = []
    run_metadata: list[pd.DataFrame] = []
    for fold_count, child_dir in completed_child_directories(sweep_dir):
        report_path = child_dir / "optimization_report.xlsx"
        overall_metrics.append(_read_sheet(report_path, REQUIRED_SHEETS[0], fold_count))
        level_metrics.append(_read_sheet(report_path, REQUIRED_SHEETS[1], fold_count))
        weight_stability.append(_read_sheet(report_path, REQUIRED_SHEETS[2], fold_count))
        run_metadata.append(_read_metadata(child_dir / "run_metadata.toml", fold_count))
    workbook_path = sweep_dir / "k_fold_comparison.xlsx"
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        pd.concat(overall_metrics, ignore_index=True).to_excel(
            writer, sheet_name="Overall Metrics by K", index=False
        )
        pd.concat(level_metrics, ignore_index=True).to_excel(
            writer, sheet_name="Per-Level Metrics by K", index=False
        )
        pd.concat(weight_stability, ignore_index=True).to_excel(
            writer, sheet_name="Fold Weight Stability by K", index=False
        )
        pd.concat(run_metadata, ignore_index=True).to_excel(
            writer, sheet_name="Run Metadata", index=False
        )
        selection_path = sweep_dir / "selection.toml"
        if selection_path.is_file():
            selection = _read_metadata(selection_path, fold_count=0).drop(columns="Fold_Count")
            selection.to_excel(writer, sheet_name="Selection", index=False)
    _format_workbook(workbook_path)
    return workbook_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--sweep-dir", type=Path, required=True)
    return parser.parse_args()


if __name__ == "__main__":
    arguments = parse_args()
    print(write_comparison_workbook(arguments.sweep_dir))