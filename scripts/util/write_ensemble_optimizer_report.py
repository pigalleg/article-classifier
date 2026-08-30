"""Create a concise Excel summary for one ensemble-weight optimization run."""

from __future__ import annotations

import argparse
from copy import copy
from pathlib import Path
from typing import cast

import pandas as pd


OVERALL_METRICS = [
    "Weighted_MAE",
    "Weighted_MSE",
    "MAE",
    "MSE",
    "Mean_Signed_Error",
    "Three_Level_Accuracy",
    "Quadratic_Weighted_Kappa",
]
LEVEL_METRICS = ["Precision", "Recall", "One_vs_Rest_Accuracy"]
LEVELS = ["Low", "Moderate", "High"]
APPROACHES = ["Optimized", "Equal_Weight"]


def _metric_matrix(
    metrics: pd.DataFrame, metric_names: list[str], index_columns: list[str] | None = None
) -> pd.DataFrame:
    index_columns = index_columns or []
    matrix = metrics.loc[metrics["Metric"].isin(metric_names)].pivot(
        index=index_columns + ["Metric"], columns="Approach", values="Value"
    )
    missing_approaches = set(APPROACHES).difference(matrix.columns)
    if missing_approaches:
        raise ValueError(f"Metrics are missing approaches: {sorted(missing_approaches)}")
    if index_columns:
        return matrix.reindex(metric_names, level="Metric")[APPROACHES]
    return matrix.reindex(metric_names)[APPROACHES]


def build_overall_metrics(metrics: pd.DataFrame) -> pd.DataFrame:
    index_columns = ["Fold"] if "Fold" in metrics.columns else []
    matrix = _metric_matrix(metrics, OVERALL_METRICS, index_columns).reset_index()
    matrix["Difference_Optimized_minus_Equal"] = (
        matrix["Optimized"] - matrix["Equal_Weight"]
    )
    return matrix


def build_level_metrics(metrics: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, float | str]] = []
    index_columns = ["Fold"] if "Fold" in metrics.columns else []
    for level in LEVELS:
        metric_names = [f"{level}_Support"] + [
            f"{level}_{metric_name}" for metric_name in LEVEL_METRICS
        ]
        matrix = _metric_matrix(metrics, metric_names, index_columns)
        fold_values = matrix.index.get_level_values("Fold").unique() if index_columns else [None]
        for fold in fold_values:
            fold_matrix = matrix.xs(fold, level="Fold") if index_columns else matrix
            optimized_support = cast(float, fold_matrix.loc[f"{level}_Support", "Optimized"])
            row: dict[str, float | str] = {
                "Level": level,
                "Support": optimized_support,
            }
            if index_columns:
                row["Fold"] = str(fold)
            for metric_name in LEVEL_METRICS:
                metric_key = f"{level}_{metric_name}"
                optimized_value = cast(float, fold_matrix.loc[metric_key, "Optimized"])
                equal_weight_value = cast(float, fold_matrix.loc[metric_key, "Equal_Weight"])
                row[f"Optimized_{metric_name}"] = optimized_value
                row[f"Equal_Weight_{metric_name}"] = equal_weight_value
                row[f"Difference_{metric_name}"] = optimized_value - equal_weight_value
            rows.append(row)
    return pd.DataFrame(rows)


def build_weight_ranges(weights: pd.DataFrame) -> pd.DataFrame:
    required_columns = {"Fold", "Grouping", "Model_Slug", "Weight"}
    missing_columns = required_columns.difference(weights.columns)
    if missing_columns:
        raise ValueError(f"Fold weights are missing columns: {sorted(missing_columns)}")
    ranges = (
        weights.groupby(["Grouping", "Model_Slug"], as_index=False)["Weight"]
        .agg(Min_Weight="min", Max_Weight="max", Mean_Weight="mean")
        .sort_values(["Grouping", "Mean_Weight"], ascending=[True, False])
    )
    ranges.insert(2, "Fold_Count", weights["Fold"].nunique())
    return ranges


def _format_workbook(workbook_path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            font = copy(cell.font)
            font.bold = True
            cell.font = font
        for column_cells in worksheet.columns:
            column_letter = column_cells[0].column_letter
            width = max(len(str(cell.value or "")) for cell in column_cells) + 2
            worksheet.column_dimensions[column_letter].width = min(width, 38)
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
    workbook.save(workbook_path)


def write_report(output_dir: Path, remove_source_csvs: bool = False) -> Path:
    metrics_path = output_dir / "metrics_summary.csv"
    fold_metrics_path = output_dir / "fold_metrics_summary.csv"
    weights_path = output_dir / "fold_weights.csv"
    if not all(path.is_file() for path in [metrics_path, fold_metrics_path, weights_path]):
        raise FileNotFoundError("Expected metrics_summary.csv, fold_metrics_summary.csv, and fold_weights.csv in the output directory.")

    metrics = pd.read_csv(metrics_path)
    fold_metrics = pd.read_csv(fold_metrics_path)
    weights = pd.read_csv(weights_path)
    workbook_path = output_dir / "optimization_report.xlsx"
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        build_overall_metrics(metrics).to_excel(writer, sheet_name="Overall Metrics", index=False)
        build_level_metrics(metrics).to_excel(writer, sheet_name="Per-Level Metrics", index=False)
        build_overall_metrics(fold_metrics).to_excel(writer, sheet_name="Overall Metrics by Fold", index=False)
        build_level_metrics(fold_metrics).to_excel(writer, sheet_name="Per-Level Metrics by Fold", index=False)
        build_weight_ranges(weights).to_excel(writer, sheet_name="Fold Weight Ranges", index=False)
    _format_workbook(workbook_path)

    if remove_source_csvs:
        metrics_path.unlink()
        fold_metrics_path.unlink()
        weights_path.unlink()
    return workbook_path


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--remove-source-csvs", action="store_true")
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    print(write_report(args.output_dir, remove_source_csvs=args.remove_source_csvs))